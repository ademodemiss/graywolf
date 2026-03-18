import argparse
import openpyxl
import os

class ExcelTool:
    def __init__(self, file_path: str):
        self.file_path = file_path

    def _load_workbook(self):
        if os.path.exists(self.file_path):
            return openpyxl.load_workbook(self.file_path)
        else:
            return openpyxl.Workbook()

    def _save_workbook(self, workbook):
        workbook.save(self.file_path)

    def read_sheet(self, sheet_name: str, start_row: int = 1, end_row: int = None, start_col: int = 1, end_col: int = None) -> list[list]:
        workbook = self._load_workbook()
        if sheet_name not in workbook.sheetnames:
            print(f"Hata: '{sheet_name}' adlı sayfa bulunamadı.")
            return []
        
        sheet = workbook[sheet_name]
        data = []
        for row_idx in range(start_row, (end_row if end_row else sheet.max_row) + 1):
            row_data = []
            for col_idx in range(start_col, (end_col if end_col else sheet.max_column) + 1):
                cell_value = sheet.cell(row=row_idx, column=col_idx).value
                row_data.append(cell_value)
            data.append(row_data)
        return data

    def write_sheet(self, sheet_name: str, data: list[list], start_row: int = 1, start_col: int = 1, overwrite: bool = False):
        workbook = self._load_workbook()
        if sheet_name not in workbook.sheetnames:
            sheet = workbook.create_sheet(sheet_name)
        else:
            sheet = workbook[sheet_name]
            if overwrite:
                # Sayfayı temizle
                if sheet.max_row > 0:
                    sheet.delete_rows(1, sheet.max_row)

        for r_idx, row_data in enumerate(data):
            for c_idx, cell_value in enumerate(row_data):
                sheet.cell(row=start_row + r_idx, column=start_col + c_idx, value=cell_value)

        self._save_workbook(workbook)
        print(f"Veriler '{sheet_name}' sayfasına başarıyla yazıldı.")

    def create_sheet(self, sheet_name: str):
        workbook = self._load_workbook()
        if sheet_name not in workbook.sheetnames:
            workbook.create_sheet(sheet_name)
            self._save_workbook(workbook)
            print(f"'{sheet_name}' sayfası başarıyla oluşturuldu.")
        else:
            print(f"Hata: '{sheet_name}' adlı sayfa zaten mevcut.")


def main():
    parser = argparse.ArgumentParser(description="GrayWolf Excel İşlemleri Aracı")
    parser.add_argument("--file", required=True, help="Excel dosya yolu")
    parser.add_argument("--action", required=True, choices=["read", "write", "create_sheet"], help="Yapılacak eylem")
    parser.add_argument("--sheet", help="Sayfa adı")
    parser.add_argument("--data", help="Yazılacak JSON veri (örn. '[[\"Başlık1\",\"Başlık2\"], [1,2]]') - write işlemi için")
    parser.add_argument("--start_row", type=int, default=1, help="Başlangıç satırı (okuma/yazma için)")
    parser.add_argument("--start_col", type=int, default=1, help="Başlangıç sütunu (okuma/yazma için)")
    parser.add_argument("--overwrite", action="store_true", help="Yazma işleminde mevcut verinin üzerine yaz")
    
    args = parser.parse_args()

    tool = ExcelTool(args.file)

    if args.action == "read":
        if not args.sheet:
            parser.error("'read' eylemi için --sheet argümanı gerekli.")
        data = tool.read_sheet(args.sheet, args.start_row, None, args.start_col, None)
        if data:
            print("Okunan Veri:")
            for row in data:
                print(row)
    elif args.action == "write":
        if not args.sheet or not args.data:
            parser.error("'write' eylemi için --sheet ve --data argümanları gerekli.")
        try:
            data_to_write = json.loads(args.data)
        except json.JSONDecodeError:
            parser.error("--data argümanı geçerli bir JSON formatında olmalı.")
        tool.write_sheet(args.sheet, data_to_write, args.start_row, args.start_col, args.overwrite)
    elif args.action == "create_sheet":
        if not args.sheet:
            parser.error("'create_sheet' eylemi için --sheet argümanı gerekli.")
        tool.create_sheet(args.sheet)


if __name__ == "__main__":
    main()
