import os
import pytest
from openpyxl import load_workbook
from tools.excel_tool import ExcelTool

@pytest.fixture
def temp_excel_file(tmp_path):
    file = tmp_path / "test.xlsx"
    yield str(file)
    if os.path.exists(file):
        os.remove(file)

def test_create_sheet(temp_excel_file):
    tool = ExcelTool(temp_excel_file)
    tool.create_sheet("TestSheet")
    workbook = load_workbook(temp_excel_file)
    assert "TestSheet" in workbook.sheetnames

def test_write_read_sheet(temp_excel_file):
    tool = ExcelTool(temp_excel_file)
    sheet_name = "DataSheet"
    data_to_write = [["Header1", "Header2"], [1, "Value2"], [3, 4]]
    tool.write_sheet(sheet_name, data_to_write)

    read_data = tool.read_sheet(sheet_name)
    assert read_data == data_to_write

def test_write_overwrite_sheet(temp_excel_file):
    tool = ExcelTool(temp_excel_file)
    sheet_name = "OverwriteSheet"
    initial_data = [["Old1", "Old2"]]
    tool.write_sheet(sheet_name, initial_data)

    new_data = [["New1", "New2"]]
    tool.write_sheet(sheet_name, new_data, overwrite=True)

    read_data = tool.read_sheet(sheet_name)
    assert read_data == new_data

def test_read_non_existent_sheet(temp_excel_file):
    tool = ExcelTool(temp_excel_file)
    data = tool.read_sheet("NonExistent")
    assert data == []
